import io
from pathlib import Path
from pypdf import PdfReader
from docx import Document

class DocumentExtractor:
    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
        """Extract text from file bytes based on file extension."""
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            return DocumentExtractor._extract_from_pdf_bytes(file_bytes)
        elif suffix == ".docx":
            return DocumentExtractor._extract_from_docx_bytes(file_bytes)
        elif suffix == ".txt":
            return file_bytes.decode("utf-8", errors="replace")
        else:
            raise ValueError(f"Unsupported file format: {suffix}")

    @staticmethod
    def extract_text_from_path(file_path: Path | str) -> str:
        """Extract text directly from a file path."""
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return DocumentExtractor._extract_from_pdf_file(path)
        elif suffix == ".docx":
            return DocumentExtractor._extract_from_docx_file(path)
        elif suffix == ".txt":
            return path.read_text(encoding="utf-8", errors="replace")
        else:
            raise ValueError(f"Unsupported file format: {suffix}")

    @staticmethod
    def _extract_from_pdf_bytes(file_bytes: bytes) -> str:
        stream = io.BytesIO(file_bytes)
        reader = PdfReader(stream)
        text_parts = []
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text.strip())
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_from_pdf_file(path: Path) -> str:
        reader = PdfReader(path)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text.strip())
        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_from_docx_bytes(file_bytes: bytes) -> str:
        stream = io.BytesIO(file_bytes)
        doc = Document(stream)
        return DocumentExtractor._parse_docx_elements(doc)

    @staticmethod
    def _extract_from_docx_file(path: Path) -> str:
        doc = Document(path)
        return DocumentExtractor._parse_docx_elements(doc)

    @staticmethod
    def _parse_docx_elements(doc: Document) -> str:
        lines = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                lines.append(txt)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    lines.append(" | ".join(row_cells))
        return "\n".join(lines)
